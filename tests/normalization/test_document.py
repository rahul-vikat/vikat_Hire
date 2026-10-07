from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter

from vikat_hire.contracts.common import InputKind
from vikat_hire.contracts.inputs import DocumentInput
from vikat_hire.contracts.normalization import ExtractionKind
from vikat_hire.normalization.document import extract_document_text


def _document(filename="resume.txt", media_type="text/plain", kind=InputKind.RESUME):
    return DocumentInput(input_id="input-001", kind=kind, filename=filename, media_type=media_type, content_hash="hash", storage_ref="storage://input")


def test_plain_text_is_deterministic_and_normalized():
    document = _document()
    first = extract_document_text(document=document, content=b"  Python Developer  \r\n\r\n Owned APIs end-to-end.\r\n", provenance_refs=("prov-1",))
    second = extract_document_text(document=document, content=b"  Python Developer  \r\n\r\n Owned APIs end-to-end.\r\n", provenance_refs=("prov-1",))
    assert first == second
    assert first[0].text == "Python Developer\n\nOwned APIs end-to-end."
    assert first[0].extraction_kind is ExtractionKind.PLAIN_TEXT


def test_plain_text_requires_utf8_and_provenance():
    with pytest.raises(ValueError, match="valid UTF-8"):
        extract_document_text(document=_document(), content=b"\xff", provenance_refs=("p",))
    with pytest.raises(ValueError, match="provenance_refs"):
        extract_document_text(document=_document(), content=b"text", provenance_refs=())


def test_empty_unsupported_and_whitespace_content_fail():
    with pytest.raises(ValueError, match="must not be empty"):
        extract_document_text(document=_document(), content=b"", provenance_refs=("p",))
    with pytest.raises(ValueError, match="empty extracted text block"):
        extract_document_text(document=_document(), content=b" \n\t", provenance_refs=("p",))
    with pytest.raises(ValueError, match="unsupported document format"):
        extract_document_text(document=_document("resume.rtf", "application/rtf"), content=b"text", provenance_refs=("p",))


def test_block_id_is_content_derived():
    document = _document()
    first = extract_document_text(document=document, content=b"Python", provenance_refs=("p1",))
    second = extract_document_text(document=document, content=b"Python", provenance_refs=("p2",))
    other = extract_document_text(document=document, content=b"Java", provenance_refs=("p1",))
    assert first[0].block_id == second[0].block_id
    assert first[0].block_id != other[0].block_id


def test_empty_pdf_fails_loudly():
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    output = BytesIO(); writer.write(output)
    with pytest.raises(ValueError, match="produced no extractable text"):
        extract_document_text(document=_document("resume.pdf", "application/pdf"), content=output.getvalue(), provenance_refs=("p",))


def test_docx_preserves_order_and_heading_context():
    buffer = BytesIO(); source = Document(); source.add_heading("Experience", level=1); source.add_paragraph("Owned APIs."); source.add_heading("Skills", level=1); source.add_paragraph("Python"); source.save(buffer)
    result = extract_document_text(document=_document("resume.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"), content=buffer.getvalue(), provenance_refs=("p",))
    assert [block.text for block in result] == ["Experience", "Owned APIs.", "Skills", "Python"]
    assert [block.section for block in result] == ["Experience", "Experience", "Skills", "Skills"]


def test_extraction_has_no_evaluation_fields():
    block = extract_document_text(document=_document(), content=b"Owned APIs.", provenance_refs=("p",))[0]
    assert not hasattr(block, "score")
    assert not hasattr(block, "candidate_level")
